# apps/reports/views.py — Views para Relatórios
# ============================================

from django.shortcuts import render
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.db.models import Count, Avg
from django.utils import timezone
from datetime import timedelta
import csv
import io

from apps.tickets.models import Ticket
from apps.accounts.models import Tecnico
from .models import RelatorioAgendado


@login_required
def relatorio_tickets(request):
    """Página de relatório de tickets com filtros"""
    tickets = Ticket.objects.select_related('solicitante', 'tecnico__usuario').all()

    # Filtros via GET
    status = request.GET.getlist('status')
    prioridade = request.GET.getlist('prioridade')
    tecnico_id = request.GET.getlist('tecnico')
    date_from = request.GET.get('date_from')
    date_to = request.GET.get('date_to')

    if status:
        tickets = tickets.filter(status__in=status)
    if prioridade:
        tickets = tickets.filter(prioridade__in=prioridade)
    if tecnico_id:
        tickets = tickets.filter(tecnico_id__in=tecnico_id)
    if date_from:
        tickets = tickets.filter(created_at__date__gte=date_from)
    if date_to:
        tickets = tickets.filter(created_at__date__lte=date_to)

    # Permissões
    if not request.user.is_superuser:
        if hasattr(request.user, 'tecnico_profile') and request.user.tecnico_profile.ativo:
            tickets = tickets.filter(models.Q(solicitante=request.user) | models.Q(tecnico=request.user.tecnico_profile))
        else:
            tickets = tickets.filter(solicitante=request.user)

    # Export
    export = request.GET.get('export')
    if export == 'csv':
        return _export_csv(tickets)
    elif export == 'excel':
        return _export_excel(tickets)

    # Agregações para gráficos
    por_status = tickets.values('status').annotate(count=Count('id'))
    por_prioridade = tickets.values('prioridade').annotate(count=Count('id'))
    por_tecnico = tickets.filter(tecnico__isnull=False).values('tecnico__usuario__first_name', 'tecnico__usuario__last_name').annotate(count=Count('id')).order_by('-count')[:10]

    # Tempo médio de resolução
    resolvidos = tickets.filter(status__in=['resolvido', 'fechado'], resolvido_em__isnull=False)
    tempo_medio = resolvidos.extra(select={'tempo_seg': "EXTRACT(EPOCH FROM (resolvido_em - created_at))"}).aggregate(avg=Avg('tempo_seg'))['avg']
    tempo_medio_str = _format_tempo(tempo_medio) if tempo_medio else '-'

    context = {
        'tickets': tickets[:100],  # Limita visualização
        'total': tickets.count(),
        'tecnicos': Tecnico.objects.filter(ativo=True).select_related('usuario'),
        'status_choices': Ticket.STATUS_CHOICES,
        'prioridade_choices': Ticket.PRIORIDADE_CHOICES,
        'por_status': por_status,
        'por_prioridade': por_prioridade,
        'por_tecnico': por_tecnico,
        'tempo_medio': tempo_medio_str,
        'current_filters': request.GET,
    }
    return render(request, 'reports/tickets.html', context)


def _export_csv(queryset):
    """Exporta tickets para CSV"""
    response = HttpResponse(content_type='text/csv; charset=utf-8')
    response['Content-Disposition'] = f'attachment; filename="tickets_{timezone.now().strftime("%Y%m%d")}.csv"'

    writer = csv.writer(response)
    writer.writerow(['ID', 'Solicitante', 'Técnico', 'Status', 'Prioridade', 'Data Ocorrido', 'Criado em', 'Resolvido em', 'Problema', 'Solução', 'Tempo Gasto'])

    for t in queryset.iterator():
        writer.writerow([
            t.pk,
            t.solicitante.get_full_name() or t.solicitante.username,
            t.tecnico.usuario.get_full_name() if t.tecnico else '',
            t.get_status_display(),
            t.get_prioridade_display(),
            t.data_ocorrido.strftime('%d/%m/%Y %H:%M'),
            t.created_at.strftime('%d/%m/%Y %H:%M'),
            t.resolvido_em.strftime('%d/%m/%Y %H:%M') if t.resolvido_em else '',
            t.problema[:200],
            t.solucao[:200],
            _format_tempo(t.tempo_gasto.total_seconds()) if t.tempo_gasto else '',
        ])

    return response


def _export_excel(queryset):
    """Exporta tickets para Excel"""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, PatternFill
    except ImportError:
        return HttpResponse('openpyxl não instalado', status=500)

    wb = Workbook()
    ws = wb.active
    ws.title = 'Tickets'

    # Header
    headers = ['ID', 'Solicitante', 'Técnico', 'Status', 'Prioridade', 'Data Ocorrido', 'Criado em', 'Resolvido em', 'Problema', 'Solução', 'Tempo Gasto']
    header_font = Font(bold=True)
    header_fill = PatternFill(start_color='3B82F6', end_color='3B82F6', fill_type='solid')
    header_font_white = Font(bold=True, color='FFFFFF')

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font_white
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')

    # Data
    for row, t in enumerate(queryset.iterator(), 2):
        ws.cell(row=row, column=1, value=t.pk)
        ws.cell(row=row, column=2, value=t.solicitante.get_full_name() or t.solicitante.username)
        ws.cell(row=row, column=3, value=t.tecnico.usuario.get_full_name() if t.tecnico else '')
        ws.cell(row=row, column=4, value=t.get_status_display())
        ws.cell(row=row, column=5, value=t.get_prioridade_display())
        ws.cell(row=row, column=6, value=t.data_ocorrido.strftime('%d/%m/%Y %H:%M'))
        ws.cell(row=row, column=7, value=t.created_at.strftime('%d/%m/%Y %H:%M'))
        ws.cell(row=row, column=8, value=t.resolvido_em.strftime('%d/%m/%Y %H:%M') if t.resolvido_em else '')
        ws.cell(row=row, column=9, value=t.problema[:500])
        ws.cell(row=row, column=10, value=t.solucao[:500])
        ws.cell(row=row, column=11, value=_format_tempo(t.tempo_gasto.total_seconds()) if t.tempo_gasto else '')

    # Ajusta largura das colunas
    for col in ws.columns:
        max_length = max(len(str(cell.value or '')) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_length + 2, 50)

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="tickets_{timezone.now().strftime("%Y%m%d")}.xlsx"'
    wb.save(response)
    return response


def _format_tempo(seconds):
    """Formata segundos em HH:MM:SS"""
    if seconds is None:
        return ''
    total = int(seconds)
    h = total // 3600
    m = (total % 3600) // 60
    s = total % 60
    return f"{h:02d}:{m:02d}:{s:02d}"


# Import models para usar no filtro
from django.db import models